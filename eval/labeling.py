# выгрузка недели на ручную разметку и импорт разметки в eval/gold_<версия>.csv,
# для дедупа — пар вакансий в eval/dedupe_gold_<версия>.csv.
# запуск из корня репозитория, export — на копии серверной базы:
#   $env:DB_PATH = "data/radar_server.db"
#   python -m eval.labeling export --gold v2 --since 2026-10-11 --until 2026-10-18
#   python -m eval.labeling import --gold v2
#   python -m eval.labeling export-pairs --gold v2 --since 2026-10-14 --until 2026-10-21
#   python -m eval.labeling import-pairs --gold v2
import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from eval.evaluate import split_keys
from pipeline import dedupe
from pipeline.dedupe import _company_key
from pipeline.normalize import normalize_dataframe
from storage.vacancies import load_raw_vacancies

DEV = "v1"  # dev-разметка: из нее подмешиваем строки для слепой переразметки
DEDUPE_DEV = "v1"
# колонки листа и их ширина
SHEET_COLUMNS = {"url": 30, "source": 10, "company": 25, "title": 50, "description": 60}
PAIR_COLUMNS = {
    "company": 18, "title_a": 40, "title_b": 40, "source_a": 10, "source_b": 10,
    "url_a": 25, "url_b": 25, "description_a": 60, "description_b": 60,
}  # fmt: skip
LABEL_COLUMNS = ["ds_role", "roles", "industries", "industry_other", "unsure", "comment"]
PAIR_LABELS = ["is_dup", "unsure", "comment"]
GOLD_PAIR_COLUMNS = ["company", "source_a", "title_a", "url_a", "source_b", "title_b", "url_b"]
MAX_DESCRIPTION = 3000


# «та же компания + то же название». без компании ключ — url,
# иначе слились бы одинаковые названия у разных скрытых работодателей
def repost_key(df):
    company = df["company"].map(_company_key)
    title = df["title"].fillna("").str.lower().str.replace(r"\W+", " ", regex=True).str.strip()
    return (company + "|" + title).where(company != "", df["url"])


# роль определяется только по названию, поэтому повтор компании + названия — тот же пример:
# из dev он дал бы утечку, внутри недели — ложно узкие бутстрэп-интервалы
def select_test(raw, dev_snapshot):
    keys = repost_key(raw)
    in_dev = raw["url"].isin(dev_snapshot["url"]) | keys.isin(repost_key(dev_snapshot))
    repost = keys.duplicated() & ~in_dev  # raw отсортирован по first_seen_at: остается первая
    print(f"за окно: {len(raw)}, пересекаются с {DEV}: {in_dev.sum()}, перепосты: {repost.sum()}")
    return raw[~in_dev & ~repost].reset_index(drop=True)


# лист без предсказаний, со ссылками: колонки данных + пустые колонки разметки
def write_sheet(df, path, columns, labels):
    df = df[list(columns)].fillna("").astype(str)
    df = df.replace(ILLEGAL_CHARACTERS_RE, "", regex=True)  # управляющие символы excel не сохранит
    for column in [c for c in df if c.startswith("description")]:
        df[column] = df[column].str.slice(0, MAX_DESCRIPTION)
    df.assign(**dict.fromkeys(labels, None)).to_excel(path, index=False, sheet_name="labels")

    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions  # сортировка фильтром не тронет шапку
    for cell, width in zip(ws[1], columns.values(), strict=False):
        ws.column_dimensions[cell.column_letter].width = width
        if cell.value.startswith("url"):
            for link in ws[cell.column_letter][1:]:
                link.hyperlink, link.style = link.value, "Hyperlink"
    wb.save(path)


def export(gold, since, until, n_recheck, seed=42):
    sheet_path = Path(f"data/eval/to_label_{gold}.xlsx")
    snapshot_path = Path(f"data/eval/snapshot_{gold}.pkl")
    if sheet_path.exists() or snapshot_path.exists():
        sys.exit(f"{sheet_path} или {snapshot_path} уже есть — не перезаписываю разметку")

    dev_snapshot = pd.read_pickle(f"data/eval/snapshot_{DEV}.pkl")
    test = select_test(load_raw_vacancies(since, until), dev_snapshot)
    if test.empty:
        sys.exit("за это окно вакансий нет — проверь DB_PATH и даты")
    test.to_pickle(snapshot_path)  # заморозка: копия базы потом перезапишется

    # поровну DS и не-DS: на 5 позитивах kappa ни о чем не скажет
    dev_gold = pd.read_csv(f"eval/gold_{DEV}.csv", encoding="utf-8-sig")
    dev_gold = dev_gold[dev_gold["unsure"].fillna(0) == 0]
    picked = dev_gold.groupby("ds_role").sample(n_recheck // 2, random_state=seed)["url"]
    recheck = dev_snapshot[dev_snapshot["url"].isin(picked)]

    sheet = pd.concat([test, recheck]).sample(frac=1, random_state=seed)
    write_sheet(sheet, sheet_path, SHEET_COLUMNS, LABEL_COLUMNS)
    print(f"на разметку {len(test) + len(recheck)}: {len(test)} новых + {len(recheck)} из {DEV}")
    print(test["source"].value_counts().to_dict(), f"-> {sheet_path}")


def cohen_kappa(a, b):
    a, b = np.asarray(a, dtype=bool), np.asarray(b, dtype=bool)
    expected = a.mean() * b.mean() + (1 - a.mean()) * (1 - b.mean())
    return ((a == b).mean() - expected) / (1 - expected)


# согласие с собой: те же вакансии из dev, размеченные вслепую второй раз
def report_agreement(dev_gold, recheck):
    df = dev_gold.merge(recheck, on="url", suffixes=("_dev", "_re"))
    df = df[(df["unsure_dev"].fillna(0) == 0) & (df["unsure_re"].fillna(0) == 0)]
    a, b = df["ds_role_dev"].astype(bool), df["ds_role_re"].astype(bool)
    both = df[a & b]
    same_roles = (both["roles_dev"].map(split_keys) == both["roles_re"].map(split_keys)).mean()
    same_ind = both["industries_dev"].map(split_keys) == both["industries_re"].map(split_keys)
    print(f"переразметка {DEV}: n = {len(df)}, ds_role совпал {(a == b).mean():.0%}")
    print(f"kappa = {cohen_kappa(a, b):.2f}, roles совпали в {same_roles:.0%} из {len(both)}")
    print(f"industries совпали в {same_ind.mean():.0%} из {len(both)}")


def import_labels(gold):
    sheet = pd.read_excel(f"data/eval/to_label_{gold}.xlsx", dtype=str)[["url", *LABEL_COLUMNS]]
    sheet = sheet.apply(lambda column: column.str.strip())
    unlabeled = sheet.loc[sheet["ds_role"].isna(), "url"]
    if len(unlabeled):
        sys.exit(f"не размечено {len(unlabeled)}: {unlabeled.head(10).tolist()}")
    for column in ["ds_role", "unsure"]:  # «1.0» из excel → 1
        sheet[column] = pd.to_numeric(sheet[column]).astype("Int64")

    dev_gold = pd.read_csv(f"eval/gold_{DEV}.csv", encoding="utf-8-sig")
    is_recheck = sheet["url"].isin(dev_gold["url"])
    test, recheck = sheet[~is_recheck], sheet[is_recheck]
    test.sort_values("url").to_csv(f"eval/gold_{gold}.csv", index=False, encoding="utf-8-sig")
    recheck.sort_values("url").to_csv(f"eval/recheck_{DEV}.csv", index=False, encoding="utf-8-sig")
    print(f"{gold}: {len(test)} вакансий, из них с DS-ролью {test['ds_role'].sum()}")
    report_agreement(dev_gold, recheck)


# === пары для дедупликации ===


# кандидаты как в dedupe v1: похожие названия, одна «голова» или общий текст ≥ 0.5.
# фильтр шире правила дедупа: любая пара, которую правило склеит, — кандидат
def is_candidate(a, b):
    overlap, jaccard = dedupe._similarity(a, b)
    nested = dedupe._titles_compatible(a["words"], b["words"])
    same_head = bool(a["head"]) and a["head"] == b["head"]
    return a["url_key"] == b["url_key"] or nested or same_head or jaccard >= 0.5 or overlap >= 0.5


# пары url → строки листа: обе вакансии рядом, в нормализованном виде, как их видит дедуп
def pair_sheet(pairs, vacancies):
    rows = vacancies.drop_duplicates("url").set_index("url")
    rows = rows.reindex(columns=["source", "company", "title", "description"])
    a = rows.reindex(pairs["url_a"]).add_suffix("_a").reset_index(drop=True)
    b = rows.reindex(pairs["url_b"]).add_suffix("_b").reset_index(drop=True)
    df = pd.concat([pairs.reset_index(drop=True), a, b], axis=1)
    return df.rename(columns={"company_a": "company"})


def export_pairs(gold, since, until, n_recheck, n_control, seed=42):
    sheet_path = Path(f"data/eval/to_label_dedupe_{gold}.xlsx")
    snapshot_path = Path(f"data/eval/dedupe_snapshot_{gold}.pkl")
    if sheet_path.exists() or snapshot_path.exists():
        sys.exit(f"{sheet_path} или {snapshot_path} уже есть — не перезаписываю разметку")

    # dev читаем до заморозки: если файла нет, снапшот не создастся и не заблокирует повтор
    dev = pd.read_csv(f"eval/dedupe_gold_{DEDUPE_DEV}.csv", encoding="utf-8-sig")
    dev_snapshot = pd.read_pickle(f"data/eval/dedupe_snapshot_{DEDUPE_DEV}.pkl")
    raw = load_raw_vacancies(since, until)
    if raw.empty:
        sys.exit("за это окно вакансий нет — проверь DB_PATH и даты")

    # все пары одной компании, как их перебирает дедуп; max_shingle_df — как в deduplicate
    df = normalize_dataframe(raw.replace(r"^\s*$", pd.NA, regex=True)).drop_duplicates("url")
    features = dedupe._build_features(df.reset_index(drop=True), max_shingle_df=5)
    pairs = pd.DataFrame(
        [(features[i]["url"], features[j]["url"], is_candidate(features[i], features[j]))
         for i, j in dedupe._company_pairs(features)],
        columns=["url_a", "url_b", "candidate"],
    )  # fmt: skip
    # контроль: случайные пары вне кандидатов — покажет дубли, которые фильтр не видит
    rest = pairs[~pairs["candidate"]]
    control = rest.sample(min(n_control, len(rest)), random_state=seed)
    test = pd.concat([pairs[pairs["candidate"]], control])
    test["sample"] = test.pop("candidate").map({True: "candidate", False: "control"})

    # слепая переразметка dev: поровну дублей и не дублей.
    # пар скрытых компаний в test быть не может — по ним переразметку было бы видно
    dev = dev[(dev["unsure"].fillna(0) == 0) & (dev["company"].map(_company_key) != "")]
    recheck = dev.groupby("is_dup").sample(n_recheck // 2, random_state=seed)[["url_a", "url_b"]]
    recheck = pair_sheet(recheck.assign(sample="recheck"), normalize_dataframe(dev_snapshot))

    sheet = pd.concat([pair_sheet(test, df), recheck]).sample(frac=1, random_state=seed)
    raw.to_pickle(snapshot_path)  # заморозка: на этом снапшоте потом считается evaluate_dedupe
    sheet[["url_a", "url_b", "sample"]].to_csv(f"data/eval/dedupe_pairs_{gold}.csv", index=False)
    write_sheet(sheet, sheet_path, PAIR_COLUMNS, PAIR_LABELS)  # sample в лист не попадает

    no_company = (df["company"].map(_company_key) == "").sum()
    print(f"вакансий: {len(df)}, из них без компании (в пары не попадают): {no_company}")
    n_candidates = pairs["candidate"].sum()
    print(f"пар одной компании: {len(pairs)}, из них кандидатов: {n_candidates}")
    print(f"на разметку {len(sheet)}: {n_candidates} кандидатов + {len(control)} контрольных")
    print(f"+ {len(recheck)} из {DEDUPE_DEV} -> {sheet_path}")


# согласие с собой по парам из dev
def report_pair_agreement(dev, recheck):
    df = dev.merge(recheck, on=["url_a", "url_b"], suffixes=("_dev", "_re"))
    df = df[(df["unsure_dev"].fillna(0) == 0) & (df["unsure_re"].fillna(0) == 0)]
    a, b = df["is_dup_dev"].astype(bool), df["is_dup_re"].astype(bool)
    print(f"переразметка {DEDUPE_DEV}: n = {len(df)}, is_dup совпал {(a == b).mean():.0%}")
    print(f"kappa = {cohen_kappa(a, b):.2f}")


def import_pairs(gold):
    sheet = pd.read_excel(f"data/eval/to_label_dedupe_{gold}.xlsx", dtype=str)
    sheet = sheet.apply(lambda column: column.str.strip())
    unlabeled = sheet.index[sheet["is_dup"].isna()] + 2  # номер строки в excel
    if len(unlabeled):
        sys.exit(f"не размечено {len(unlabeled)} пар, строки: {unlabeled[:10].tolist()}")
    for column in ["is_dup", "unsure"]:
        sheet[column] = pd.to_numeric(sheet[column]).astype("Int64")

    samples = pd.read_csv(f"data/eval/dedupe_pairs_{gold}.csv")
    sheet = sheet.merge(samples, on=["url_a", "url_b"], validate="one_to_one")
    columns = [*GOLD_PAIR_COLUMNS, *PAIR_LABELS, "sample"]
    test, recheck = sheet[sheet["sample"] != "recheck"], sheet[sheet["sample"] == "recheck"]
    test[columns].sort_values(["sample", "company"]).to_csv(
        f"eval/dedupe_gold_{gold}.csv", index=False, encoding="utf-8-sig"
    )
    recheck[columns[:-1]].sort_values("company").to_csv(
        f"eval/dedupe_recheck_{DEDUPE_DEV}.csv", index=False, encoding="utf-8-sig"
    )

    for name, part in test.groupby("sample"):
        dups, unsure = part["is_dup"].sum(), part["unsure"].sum()
        print(f"{name}: {len(part)} пар, дублей {dups}, unsure {unsure}")
    dev = pd.read_csv(f"eval/dedupe_gold_{DEDUPE_DEV}.csv", encoding="utf-8-sig")
    report_pair_agreement(dev, recheck)


def main():
    parser = argparse.ArgumentParser()
    actions = ["export", "import", "export-pairs", "import-pairs"]
    parser.add_argument("action", choices=actions)
    parser.add_argument("--gold", default="v2", help="версия разметки")
    parser.add_argument("--since", help="начало окна по first_seen_at, YYYY-MM-DD")
    parser.add_argument("--until", help="конец окна, не включая")
    parser.add_argument("--recheck", type=int, help="из dev на переразметку: 50 вакансий / 30 пар")
    parser.add_argument("--control", type=int, default=100, help="контрольных пар вне кандидатов")
    args = parser.parse_args()

    if args.action.startswith("export") and not (args.since and args.until):
        parser.error("для export нужны --since и --until")
    if args.action == "export":
        export(args.gold, args.since, args.until, args.recheck or 50)
    elif args.action == "export-pairs":
        export_pairs(args.gold, args.since, args.until, args.recheck or 30, args.control)
    elif args.action == "import":
        import_labels(args.gold)
    else:
        import_pairs(args.gold)


if __name__ == "__main__":
    main()
